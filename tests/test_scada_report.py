"""Logs against synthetic SCADA (lost energy, agreement, clock check, code catalogue) and the KPI pack files."""
import numpy as np
import pandas as pd
import pytest

from wind_om.analysis import analyse
from wind_om.synthetic import mapping_frame, synthetic_log

wd = pytest.importorskip("wind_downtime")  # the "scada" extra

RATED = 2000.0


def _curve(ws):
    p = RATED / (1 + np.exp(-(np.asarray(ws) - 8.0) / 1.2))
    return np.where((ws < 3.0) | (ws > 25.0), 0.0, p)


def make_scada(days=30, n=5, seed=3):
    """10-minute SCADA of n turbines with logged stops painted in: (long table, log, truth kWh)."""
    rng = np.random.default_rng(seed)
    t = pd.date_range("2024-03-01", periods=days * 144, freq="10min", tz="UTC")
    base = np.clip(8.5 + 4 * np.sin(np.arange(len(t)) / 250) + np.cumsum(rng.normal(0, 0.04, len(t))), 3.5, 20)
    ws = base[:, None] * (1 + rng.normal(0, 0.03, (len(t), n)))
    power = _curve(ws) * (1 + rng.normal(0, 0.02, ws.shape))
    pitch = np.clip((ws - 11.5) * 2.5, 0, 30) + rng.normal(0, 0.3, ws.shape)
    rpm = np.clip(6 + 0.9 * ws, 0, 15) + rng.normal(0, 0.1, ws.shape)
    rows, truth = [], 0.0
    for k, (i, a, b, code) in enumerate([(0, 600, 640, "1001"), (1, 1500, 1512, "2002"), (2, 2600, 2700, "1001"),
                                         (3, 3300, 3336, "2002"), (4, 900, 930, "3003")]):
        truth += _curve(ws[a:b, i]).sum() / 6
        power[a:b, i], pitch[a:b, i], rpm[a:b, i] = -2.0, 90.0, 0.0
        rows.append((f"WTG{i + 1:02d}", t[a], t[b], code, "fault", "forced_outage", True))
    long = pd.DataFrame({"timestamp": np.repeat(t, n), "turbine_id": np.tile([f"WTG{i + 1:02d}" for i in range(n)],
                                                                          len(t)),
                         "power": power.ravel(), "wind_speed": ws.ravel(), "rotor_speed": rpm.ravel(),
                         "pitch": pitch.ravel(), "nacelle_dir": 200.0, "ambient_temp": 10.0})
    log = pd.DataFrame(rows, columns=["turbine_id", "start", "end", "code", "description", "category", "stop"])
    return long, log, truth


@pytest.fixture(scope="module")
def setup():
    from wind_downtime.schema import to_fleet

    from wind_om.scada import SCADA
    from wind_om.schema import validate

    long, log, truth = make_scada()
    sc = SCADA.build(to_fleet(long, rated_kw=RATED))
    a = analyse(validate(log), mapping_frame())
    return a, sc, truth, log


def test_lost_energy_of_logged_stops(setup):
    from wind_om.scada import attribute_lost

    a, sc, truth, _ = setup
    seg = attribute_lost(a.seg, sc)
    assert seg["lost_kwh"].sum() == pytest.approx(truth, rel=0.08)


def test_agreement_and_clock_check(setup):
    from wind_om.scada import agreement, clock_check

    a, sc, _, _ = setup
    ov, mon = agreement(a.seg, sc)
    assert ov["f1"] > 0.95
    assert mon["log_h"].sum() == pytest.approx(ov["log_hours"], rel=0.02)
    # a log written in local summer time but read as UTC is two hours late: the check finds the shift
    late = a.seg.assign(start=a.seg["start"] + pd.Timedelta("2h"), end=a.seg["end"] + pd.Timedelta("2h"))
    cc = clock_check(late, sc, max_shift_h=3)
    assert cc.loc[cc["f1"].idxmax(), "shift_h"] == pytest.approx(-2.0)


def test_code_stop_share(setup):
    from wind_om.scada import code_stop_share
    from wind_om.schema import validate

    _, sc, _, log = setup
    info = log.assign(code="9", description="warning", category="information", stop=False,
                      start=log["start"] + pd.Timedelta("2D"), end=log["end"] + pd.Timedelta("2D"))
    ev = validate(pd.concat([log, info], ignore_index=True))
    cat = code_stop_share(ev, sc, min_intervals=6).set_index("code")
    assert cat.loc["1001", "stop_share"] > 0.9 and cat.loc["9", "stop_share"] < 0.1


def test_kpi_pack_files(tmp_path):
    from openpyxl import load_workbook

    from wind_om.report import kpi_pack, write_excel, write_html

    log, _ = synthetic_log(n_turbines=8, years=1.5, process="hpp", mtbf_h=300, start="2025-06-01",
                           repeat_prob=0.2, cascade_prob=0.5, onsite_prob=0.5, seed=12)
    a = analyse(log, mapping_frame())
    pack = kpi_pack(a, "2026-09", tz="Europe/Warsaw", farm="test farm")
    assert pack.summary["Failures (month)"] == len(a.fail[(a.fail["start"] >= pd.Timestamp("2026-08-31 22:00",
                                                                                           tz="UTC"))
                                                         & (a.fail["start"] < pd.Timestamp("2026-09-30 22:00",
                                                                                           tz="UTC"))])
    x = write_excel(pack, tmp_path / "kpi.xlsx")
    h = write_html(pack, tmp_path / "kpi.html")
    wb = load_workbook(x)
    assert {"Summary", "Pareto codes", "MTBF turbines", "Bad actors", "Repeats", "Failures"} <= set(wb.sheetnames)
    assert len(wb["Pareto codes"].conditional_formatting) >= 2
    assert len(wb["MTBF turbines"].conditional_formatting) >= 2
    text = h.read_text(encoding="utf-8")
    assert "<svg" in text and "Pareto by code" in text and "test farm" in text
    with pytest.raises(ValueError):
        kpi_pack(a, "2031-01", tz="Europe/Warsaw")


def test_cli_demo(tmp_path):
    from wind_om.cli import main

    assert main(["demo", "--out", str(tmp_path)]) == 0
    assert (tmp_path / "kpi_2026-10.xlsx").exists() and (tmp_path / "kpi_2026-10.html").exists()
    assert main(["report", str(tmp_path / "demo_log.csv"), "--config", str(tmp_path / "demo_config.yaml"),
                 "--mapping", str(tmp_path / "demo_mapping.csv"), "--month", "2026-03", "--out",
                 str(tmp_path / "r")]) == 0
    assert main(["template", str(tmp_path / "demo_log.csv"), "--config", str(tmp_path / "demo_config.yaml"),
                 "--out", str(tmp_path / "codes.csv")]) == 0
    assert len(pd.read_csv(tmp_path / "codes.csv")) >= 6
